"""
Excel workbook with LIVE formulas: Inputs -> Scorecard -> Pricing.

The workbook mirrors the Python logic cell by cell (ratio, interpolated factor
score, weighted score, notching, log-interpolated spread) so a reviewer can
change a financial in Inputs and watch the rating and the rate move.
`verify_against_python()` recalculates the file with LibreOffice and checks the
key cells against the Python model before the file is delivered.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .scorecard import BUCKETS, NOTCHES, Financials, score, with_implicit_support
from .curves import SECURED_ADJ_BPS, SUBORDINATED_ADJ_BPS, price_loan, rating_grid

BLUE_INPUT = Font(color="0000FF")
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

FACTOR_LABELS = {
    "ebitda_margin": "EBITDA margin",
    "debt_to_ebitda": "Debt / EBITDA",
    "ebitda_to_interest": "EBITDA / interest",
    "ffo_to_debt": "FFO / debt",
    "revenue_eur_m": "Revenue (EUR m)",
}
RATIO_FORMULA = {   # in terms of Inputs cells
    "ebitda_margin": "=Inputs!$B$3/Inputs!$B$2",
    "debt_to_ebitda": "=Inputs!$B$6/Inputs!$B$3",
    "ebitda_to_interest": "=Inputs!$B$3/Inputs!$B$4",
    "ffo_to_debt": "=(Inputs!$B$3-Inputs!$B$4-Inputs!$B$5)/Inputs!$B$6",
    "revenue_eur_m": "=Inputs!$B$2",
}
TENORS = (1, 2, 3, 5, 7, 10)


def write_workbook(case: dict, grid: dict, path: str | Path, curves=None, rates=None) -> Path:
    fin = Financials(**{k: case[k] for k in ("revenue", "ebitda", "interest_expense", "tax_paid", "total_debt")})
    wb = Workbook()

    # ---- Inputs
    inp = wb.active
    inp.title = "Inputs"
    rows = [("Revenue (EUR m)", fin.revenue, "#,##0.0"), ("EBITDA", fin.ebitda, "#,##0.0"),
            ("Interest expense", fin.interest_expense, "#,##0.0"), ("Tax paid", fin.tax_paid, "#,##0.0"),
            ("Total debt", fin.total_debt, "#,##0.0"), ("Group rating", case["group_rating"], "@"),
            ("Max implicit-support notches", case["implicit_support_notches"], "0"),
            ("Currency", case["currency"], "@"), ("Tenor (years)", case["tenor_years"], "0.0"),
            ("Principal (EUR m)", case["principal"], "#,##0.0"),
            ("Secured (1/0)", int(case["secured"]), "0"), ("Subordinated (1/0)", int(case["subordinated"]), "0")]
    inp["A1"], inp["B1"] = "Input", "Value"
    inp["A1"].font = inp["B1"].font = BOLD
    for i, (label, v, fmt) in enumerate(rows, start=2):
        inp.cell(i, 1, label)
        c = inp.cell(i, 2, v)
        c.font, c.number_format, c.border = BLUE_INPUT, fmt, BOX
    inp.column_dimensions["A"].width = 30
    inp.column_dimensions["B"].width = 14

    # ---- Scorecard: thresholds, ratio, score (interpolated), weight
    sc = wb.create_sheet("Scorecard")
    head = ["Factor", "Ratio", "Weight", "Higher is better", "Score", "Weighted"] + BUCKETS
    for j, h in enumerate(head, start=1):
        c = sc.cell(1, j, h); c.font = BOLD; c.fill = HEAD_FILL
    r = 2
    for name, spec in grid["factors"].items():
        sc.cell(r, 1, FACTOR_LABELS[name])
        sc.cell(r, 2, RATIO_FORMULA[name]).number_format = "0.000"
        sc.cell(r, 3, spec["weight"]).number_format = "0%"
        sc.cell(r, 3).font = BLUE_INPUT
        sc.cell(r, 4, 1 if spec["higher_is_better"] else 0).font = BLUE_INPUT
        for j, t in enumerate(spec["thresholds"]):
            c = sc.cell(r, 7 + j, t); c.font = BLUE_INPUT; c.number_format = "0.00"
        rng = f"$G${r}:$M${r}"
        x = f"$B${r}"
        # higher_is_better: thresholds descending -> MATCH(x, rng, -1) gives k with T[k] >= x > T[k+1]
        # lower_is_better : thresholds ascending  -> MATCH(x, rng,  1) gives k with T[k] <= x < T[k+1]
        hib = (f"IF({x}>=$G${r},1,IF({x}<=$M${r},7,"
               f"MATCH({x},{rng},-1)+(INDEX({rng},MATCH({x},{rng},-1))-{x})"
               f"/(INDEX({rng},MATCH({x},{rng},-1))-INDEX({rng},MATCH({x},{rng},-1)+1))))")
        lib = (f"IF({x}<=$G${r},1,IF({x}>=$M${r},7,"
               f"MATCH({x},{rng},1)+({x}-INDEX({rng},MATCH({x},{rng},1)))"
               f"/(INDEX({rng},MATCH({x},{rng},1)+1)-INDEX({rng},MATCH({x},{rng},1)))))")
        sc.cell(r, 5, f"=IF($D${r}=1,{hib},{lib})").number_format = "0.00"
        sc.cell(r, 6, f"=$C${r}*$E${r}").number_format = "0.000"
        r += 1
    last = r - 1
    sc.cell(r + 1, 1, "Weighted score").font = BOLD
    sc.cell(r + 1, 6, f"=SUM(F2:F{last})").number_format = "0.00"
    sc.cell(r + 2, 1, "Stand-alone rating").font = BOLD
    # score -> notch: 1-based index 3 + round((score-2)*3), clamped to the list (AA is row 3)
    sc.cell(r + 2, 6, f'=INDEX(Notches!$A$1:$A$21,MIN(21,MAX(1,3+ROUND((F{r + 1}-2)*3,0))))')
    sc.cell(r + 3, 1, "Implicit-support lift (notches)").font = BOLD
    sc.cell(r + 3, 6, f"=MIN(MAX(MATCH(F{r + 2},Notches!$A$1:$A$21,0)-MATCH(Inputs!$B$7,Notches!$A$1:$A$21,0),0),Inputs!$B$8)")
    sc.cell(r + 4, 1, "Supported rating").font = BOLD
    sc.cell(r + 4, 6, f"=INDEX(Notches!$A$1:$A$21,MATCH(F{r + 2},Notches!$A$1:$A$21,0)-F{r + 3})")
    score_row, sa_row, lift_row, sup_row = r + 1, r + 2, r + 3, r + 4
    sc.column_dimensions["A"].width = 30
    for col in "BCDEF":
        sc.column_dimensions[col].width = 12

    # ---- Notches (lookup list)
    nt = wb.create_sheet("Notches")
    for i, n in enumerate(NOTCHES, start=1):
        nt.cell(i, 1, n)

    # ---- Spreads: notch x tenor grid (values from the Python curve model)
    sp = wb.create_sheet("Spreads")
    g = rating_grid(TENORS, curves)
    sp.cell(1, 1, "Rating \\ tenor").font = BOLD
    for j, t in enumerate(TENORS, start=2):
        c = sp.cell(1, j, t); c.font = BOLD; c.fill = HEAD_FILL
    for i, (rt, row) in enumerate(g.iterrows(), start=2):
        sp.cell(i, 1, rt).font = BOLD
        for j, t in enumerate(TENORS, start=2):
            sp.cell(i, j, round(float(row[t]), 1)).number_format = "0.0"
    n_rows = len(g)
    tenor_hdr = f"Spreads!$B$1:${get_column_letter(1 + len(TENORS))}$1"
    rating_col = f"Spreads!$A$2:$A${1 + n_rows}"
    grid_rng = f"Spreads!$B$2:${get_column_letter(1 + len(TENORS))}${1 + n_rows}"

    # ---- Base rates
    br = wb.create_sheet("BaseRates")
    from .curves import load_base_rates
    rates = rates or load_base_rates()
    br.cell(1, 1, "Currency \\ tenor").font = BOLD
    for j, t in enumerate(TENORS, start=2):
        c = br.cell(1, j, t); c.font = BOLD; c.fill = HEAD_FILL
    for i, (ccy, s) in enumerate(rates.items(), start=2):
        br.cell(i, 1, ccy).font = BOLD
        for j, t in enumerate(TENORS, start=2):
            c = br.cell(i, j, float(s.get(float(t)))); c.font = BLUE_INPUT; c.number_format = "0.00%"
    br_hdr = f"BaseRates!$B$1:${get_column_letter(1 + len(TENORS))}$1"
    br_rng = f"BaseRates!$B$2:${get_column_letter(1 + len(TENORS))}${1 + len(rates)}"
    br_ccy = f"BaseRates!$A$2:$A${1 + len(rates)}"

    # ---- Pricing
    pr = wb.create_sheet("Pricing")
    t = "Inputs!$B$10"
    k = f"MATCH({t},{tenor_hdr},1)"                       # bracketing lower tenor column
    rr = f"MATCH(Scorecard!$F${sup_row},{rating_col},0)"  # rating row
    cr = f"MATCH(Inputs!$B$9,{br_ccy},0)"                 # currency row
    def loginterp(rng, rowm):
        lo = f"INDEX({rng},{rowm},{k})"; hi = f"INDEX({rng},{rowm},{k}+1)"
        tlo = f"INDEX({tenor_hdr},1,{k})"; thi = f"INDEX({tenor_hdr},1,{k}+1)"
        return (f"IF({t}>=INDEX({tenor_hdr},1,{len(TENORS)}),INDEX({rng},{rowm},{len(TENORS)}),"
                f"IF({t}={tlo},{lo},EXP(LN({lo})+(LN({t})-LN({tlo}))/(LN({thi})-LN({tlo}))*(LN({hi})-LN({lo})))))")
    def lininterp(rng, rowm):
        lo = f"INDEX({rng},{rowm},{k})"; hi = f"INDEX({rng},{rowm},{k}+1)"
        tlo = f"INDEX({tenor_hdr},1,{k})"; thi = f"INDEX({tenor_hdr},1,{k}+1)"
        return (f"IF({t}>=INDEX({tenor_hdr},1,{len(TENORS)}),INDEX({rng},{rowm},{len(TENORS)}),"
                f"IF({t}={tlo},{lo},{lo}+(LN({t})-LN({tlo}))/(LN({thi})-LN({tlo}))*({hi}-{lo})))")
    lines = [
        ("Supported rating", f"=Scorecard!$F${sup_row}", "@"),
        ("Base rate", "=" + lininterp(br_rng, cr), "0.000%"),
        ("Credit spread (bps)", "=" + loginterp(grid_rng, rr), "0.0"),
        ("Security / subordination adj. (bps)", f"=Inputs!$B$12*({SECURED_ADJ_BPS})+Inputs!$B$13*({SUBORDINATED_ADJ_BPS})", "0"),
        ("All-in interest rate", "=B3+(B4+B5)/10000", "0.000%"),
        ("Annual interest (EUR m)", "=B6*Inputs!$B$11", "#,##0.000"),
    ]
    pr["A1"], pr["B1"] = "Item", "Value"
    pr["A1"].font = pr["B1"].font = BOLD
    for i, (label, f, fmt) in enumerate(lines, start=2):
        pr.cell(i, 1, label)
        c = pr.cell(i, 2, f); c.number_format = fmt; c.border = BOX
    pr.column_dimensions["A"].width = 36
    pr.column_dimensions["B"].width = 14

    path = Path(path)
    wb.save(path)
    return path


def recalculate(path: Path) -> Path | None:
    """Recalculate with LibreOffice (if installed) so cached values exist. Returns the recalculated copy."""
    if shutil.which("soffice") is None:
        return None
    tmp = Path(tempfile.mkdtemp())
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp), str(path)],
                   check=True, capture_output=True, timeout=120)
    return tmp / path.name


def verify_against_python(path: Path, case: dict, grid: dict, tol_rate=1e-6, tol_bps=0.15) -> dict:
    """Compare the recalculated workbook with the Python model. Raises on mismatch."""
    rc = recalculate(path)
    if rc is None:
        return {"checked": False}
    wb = load_workbook(rc, data_only=True)
    fin = Financials(**{k: case[k] for k in ("revenue", "ebitda", "interest_expense", "tax_paid", "total_debt")})
    res = score(fin, grid)
    sup, _ = with_implicit_support(res["standalone_rating"], case["group_rating"], case["implicit_support_notches"])
    px = price_loan(sup, case["currency"], case["tenor_years"], case["secured"], case["subordinated"])
    sc = wb["Scorecard"]; pr = wb["Pricing"]
    n = len(grid["factors"])
    xl = {"score": sc.cell(n + 3, 6).value, "standalone": sc.cell(n + 4, 6).value, "supported": sc.cell(n + 6, 6).value,
          "spread_bps": pr["B4"].value, "all_in": pr["B6"].value}
    checks = [
        ("score", abs(xl["score"] - res["score"]) < 1e-6),
        ("standalone rating", xl["standalone"] == res["standalone_rating"]),
        ("supported rating", xl["supported"] == sup),
        ("spread", abs(xl["spread_bps"] - px["spread_bps"]) < tol_bps),
        ("all-in rate", abs(xl["all_in"] - px["all_in_rate"]) < tol_rate + tol_bps / 1e4),
    ]
    bad = [name for name, ok in checks if not ok]
    if bad:
        raise AssertionError(f"Excel and Python disagree on: {bad}; excel={xl}")
    shutil.copy(rc, path)   # ship the recalculated file so values show without a recalc
    return {"checked": True, "excel": xl}
